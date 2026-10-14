"""Executive readout deck (Section 6.9 #1): 14 slides, names and colours from theme tokens."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from altitude.core import theme as theme_mod

from .engine import DiagnosticResult

# Layout geometry in EMU-friendly inches (presentation layout, not thresholds).
SLIDE_W, SLIDE_H = 13.333, 7.5


def _rgb(hex_token: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_token.lstrip("#"))


class _Deck:
    def __init__(self, result: DiagnosticResult, theme: str | None):
        from pptx import Presentation
        from pptx.util import Inches

        self.r = result
        self.t = lambda k: theme_mod.token(k, theme)
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        dt = result.context.window_end
        self.prs.core_properties.title = f"{self.t('PACKAGE_D_NAME')} — {result.snapshot_id}"
        self.prs.core_properties.subject = f"snapshot_id={result.snapshot_id}; rubric_version={result.rubric_version}"
        self.prs.core_properties.created = dt
        self.prs.core_properties.modified = dt
        self.prs.core_properties.last_modified_by = self.t("PRODUCT_NAME")
        self.prs.core_properties.author = self.t("COMPANY_NAME")

    # ------------------------------------------------------------------ primitives
    def text(self, slide, x, y, w, h, text: str, size: int = 14, bold: bool = False, color: str = "COLOR_TEXT",
             align: str | None = None):
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = _rgb(self.t(color))
            if align == "center":
                p.alignment = PP_ALIGN.CENTER
        return box

    def slide(self, title: str, subtitle: str | None = None):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches

        s = self.prs.slides.add_slide(self.blank)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, Inches(1.0))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(self.t("COLOR_PRIMARY"))
        bar.line.fill.background()
        self.text(s, 0.5, 0.18, 12.3, 0.7, title, size=26, bold=True, color="COLOR_SURFACE")
        if subtitle:
            self.text(s, 0.5, 1.1, 12.3, 0.5, subtitle, size=14, color="COLOR_PRIMARY")
        self.text(s, 0.5, 7.0, 12.3, 0.35,
                  f"{self.t('PRODUCT_NAME')} · {self.t('PACKAGE_D_NAME')} · snapshot {self.r.snapshot_id} · "
                  f"rubric {self.r.rubric_version} · {len(self.prs.slides)}", size=9, color="COLOR_PRIMARY")
        return s

    def table(self, slide, x, y, w, rows: list[list[Any]], col_widths: list[float] | None = None, size: int = 11):
        from pptx.util import Inches, Pt

        n_rows, n_cols = len(rows), len(rows[0])
        h = 0.32 * n_rows
        shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(h))
        tbl = shape.table
        if col_widths:
            for i, cw in enumerate(col_widths):
                tbl.columns[i].width = Inches(cw)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = tbl.cell(ri, ci)
                cell.text = "" if val is None else str(val)
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(size)
                para.font.bold = ri == 0
                if ri == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(self.t("COLOR_PRIMARY"))
                    para.font.color.rgb = _rgb(self.t("COLOR_SURFACE"))
        return tbl

    def bar_chart(self, slide, x, y, w, h, categories: list[str], series: dict[str, list[float]], color: str = "COLOR_ACCENT"):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.util import Inches

        data = CategoryChartData()
        data.categories = categories
        for name, values in series.items():
            data.add_series(name, values)
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
        chart = gf.chart
        chart.has_legend = len(series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        palette = [color, "COLOR_PRIMARY", "COLOR_WARN"]
        for i, s in enumerate(chart.series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = _rgb(self.t(palette[i % len(palette)]))
        return chart


def _usd(v: float | int | None) -> str:
    return "—" if v is None else f"${v:,.0f}"


def _val(s: dict) -> str:
    v = s["value"]["annualised_usd"]
    return _usd(v["p50"]) if v else "risk"


def _risk(s: dict) -> str:
    r = s["risk_reduction"]
    return f"{r['score']}→{r['residual_score']}" if r else "—"


def build_deck(result: DiagnosticResult, path: str | Path, recs: list[dict] | None = None, theme: str | None = None) -> int:
    """Write the executive readout; returns the number of slides."""
    recs = recs if recs is not None else result.suggestions
    d = _Deck(result, theme)
    sc = result.scorecard
    ctx = result.context
    phi = sc["platform_health_index"]
    account = ctx.account.get("ACCOUNT_NAME") or ctx.manifest.get("metadata", {}).get("account_name") or ""
    total_p50 = sum((s["value"]["annualised_usd"] or {}).get("p50", 0) for s in recs)
    total_cons = sum((s["value"]["annualised_usd"] or {}).get("conservative", 0) for s in recs)
    by_q = {q: [s for s in recs if s["priority"]["quadrant"] == q] for q in ("quick_win", "strategic", "hygiene", "defer")}
    review = [s for s in recs if result.detail(s["id"]).needs_review]

    # 1 title
    s = d.slide(f"{d.t('PACKAGE_D_NAME')} — Executive readout")
    d.text(s, 0.8, 2.2, 11.5, 1.0, account, size=40, bold=True, color="COLOR_PRIMARY")
    d.text(s, 0.8, 3.4, 11.5, 1.6,
           f"Evidence window {ctx.window_start.date()} to {ctx.window_end.date()} · tier {result.tier}\n"
           f"Sealed snapshot {result.snapshot_id} · rubric {result.rubric_version}\n{d.t('COMPANY_NAME')}", size=18)

    # 2 executive summary
    s = d.slide("Executive summary", "Five questions: waste, exposure, fragility, trust in reports, AI readiness")
    tiles = [(f"{phi:.0f}" if phi is not None else "—", d.t("HEALTH_INDEX_NAME")), (_usd(total_p50), "identified value / yr (p50)"),
             (_usd(total_cons), "conservative / yr"), (str(len(recs)), "suggestions"), (str(len(by_q["quick_win"])), "quick wins")]
    for i, (v, label) in enumerate(tiles):
        d.text(s, 0.5 + i * 2.5, 1.9, 2.4, 0.8, v, size=30, bold=True, color="COLOR_ACCENT", align="center")
        d.text(s, 0.5 + i * 2.5, 2.7, 2.4, 0.6, label, size=12, align="center")
    lines = [f"{x['priority']['rank']}. {x['title']} ({_val(x)})" for x in recs[:5]]
    d.text(s, 0.8, 3.7, 11.8, 3.0, "Top priorities\n" + "\n".join(lines), size=16)

    # 3 PHI + domain scorecard table
    s = d.slide(f"{d.t('HEALTH_INDEX_NAME')}: {phi if phi is not None else '—'} / 100", "Scorecard by domain (1 Initial – 5 Optimised)")
    rows = [["Domain", "Name", "Score", "Maturity", "Checks failing"]]
    rows += [[x["id"], x["name"], x["score"] if x["score"] else "—", x["label"] or "not in scope", x["checks_failed"]]
             for x in sc["domains"]]
    d.table(s, 0.8, 1.7, 11.5, rows, [1.2, 5.0, 1.2, 2.3, 1.8])

    # 4 domain chart
    s = d.slide("Domain scores", "Weighted check outcomes per domain (rubric bands)")
    scored = [x for x in sc["domains"] if x["score"]]
    d.bar_chart(s, 0.8, 1.6, 11.5, 5.2, [x["id"] for x in scored], {"Score (1–5)": [float(x["score"]) for x in scored]})

    # 5 WAF pillars
    s = d.slide("Well-Architected pillar roll-up", "Same five pillars as the free WAF assessment, scored from evidence")
    rows = [["Pillar", "Score", "Maturity", "Index"]] + [[p["name"], p["score"] or "—", p["label"] or "—", p["index"]]
                                                          for p in sc["pillars"]]
    d.table(s, 0.8, 1.7, 6.2, rows, [2.6, 1.0, 1.4, 1.2])
    d.bar_chart(s, 7.3, 1.6, 5.5, 4.8, [p["pillar"].replace("_", " ") for p in sc["pillars"]],
                {"Index": [float(p["index"] or 0) for p in sc["pillars"]]}, color="COLOR_PRIMARY")

    # 6 top 10
    s = d.slide("Top 10 suggestions", "Ranked by value × risk ÷ effort × confidence")
    rows = [["#", "Suggestion", "Value / yr", "Risk", "Effort", "Conf."]]
    rows += [[x["priority"]["rank"], x["title"][:70], _val(x), _risk(x), x["effort"]["size"], f"{x['confidence']:.2f}"]
             for x in recs[:10]]
    d.table(s, 0.5, 1.6, 12.3, rows, [0.5, 7.0, 1.4, 1.1, 1.0, 1.0], size=10)

    # 7 quick wins
    s = d.slide("Quick wins — first 30 days", "High value, low effort")
    rows = [["Suggestion", "Value / yr", "Effort (h)", "Autonomy"]]
    rows += [[x["title"][:75], _val(x), x["effort"]["hours_p50"], x["remediation"]["autonomy"]] for x in by_q["quick_win"][:12]]
    d.table(s, 0.5, 1.6, 12.3, rows if len(rows) > 1 else rows + [["—", "", "", ""]], [7.8, 1.7, 1.4, 1.4], size=10)

    # 8 strategic
    s = d.slide("Strategic — 90-day plan", "High value, higher effort")
    rows = [["Suggestion", "Value / yr", "Risk", "Effort (h)"]]
    rows += [[x["title"][:75], _val(x), _risk(x), x["effort"]["hours_p50"]] for x in by_q["strategic"][:12]]
    d.table(s, 0.5, 1.6, 12.3, rows if len(rows) > 1 else rows + [["—", "", "", ""]], [7.8, 1.7, 1.4, 1.4], size=10)

    # 9 savings model
    s = d.slide("Savings model", f"Contract price {_usd(ctx.price_per_credit) if ctx.price_per_credit >= 1 else ctx.price_per_credit} "
                                  f"per credit ({ctx.price_source.replace('_', ' ')}); annualised from a {int(ctx.window_days)}-day window")
    doms = sorted({x["domain"] for x in recs if x["value"]["annualised_usd"]}, key=lambda z: int(z[1:]))
    p50 = [float(sum((x["value"]["annualised_usd"] or {}).get("p50", 0) for x in recs if x["domain"] == dm)) for dm in doms]
    cons = [float(sum((x["value"]["annualised_usd"] or {}).get("conservative", 0) for x in recs if x["domain"] == dm)) for dm in doms]
    if doms:
        d.bar_chart(s, 0.8, 1.7, 8.0, 5.0, doms, {"p50": p50, "conservative": cons})
    d.text(s, 9.1, 1.9, 3.8, 4.5, f"Total p50: {_usd(total_p50)}\nConservative: {_usd(total_cons)}\n\n"
                                  "Every value carries its method and evidence IDs; assumptions are listed in the "
                                  "savings & risk model workbook.", size=14)

    # 10 risk model
    s = d.slide("Risk model", "Likelihood × impact (1–5 each) and residual after remediation")
    risky = sorted([x for x in recs if x["risk_reduction"]], key=lambda x: -x["risk_reduction"]["score"])[:12]
    rows = [["Suggestion", "L", "I", "Score", "Residual"]]
    rows += [[x["title"][:80], x["risk_reduction"]["likelihood"], x["risk_reduction"]["impact"], x["risk_reduction"]["score"],
              x["risk_reduction"]["residual_score"]] for x in risky]
    d.table(s, 0.5, 1.6, 12.3, rows if len(rows) > 1 else rows + [["—", "", "", "", ""]], [8.7, 0.8, 0.8, 1.0, 1.0], size=10)

    # 11 remediation & governance
    s = d.slide("Remediation-as-code and review", "Drafts with rollback; humans approve every change")
    autonomy = {}
    for x in recs:
        autonomy[x["remediation"]["autonomy"]] = autonomy.get(x["remediation"]["autonomy"], 0) + 1
    rows = [["Measure", "Count"], ["Remediation drafts (each with rollback)", len(recs)]]
    rows += [[f"Autonomy {k}", v] for k, v in sorted(autonomy.items())]
    rows += [["Mandatory review (confidence below floor)", len(review)],
             ["Sampled for review", sum(1 for x in recs if result.detail(x["id"]).review_reason == "sampled")]]
    d.table(s, 0.8, 1.7, 7.0, rows, [5.0, 2.0])

    # 12 roadmap
    s = d.slide("Roadmap", "30 / 90 days / backlog")
    cols = [("First 30 days", by_q["quick_win"]), ("Sprint batch", by_q["hygiene"]), ("90-day plan", by_q["strategic"]),
            ("Backlog", by_q["defer"])]
    for i, (name, items) in enumerate(cols):
        d.text(s, 0.5 + i * 3.15, 1.6, 3.0, 0.5, f"{name} ({len(items)})", size=16, bold=True, color="COLOR_PRIMARY")
        d.text(s, 0.5 + i * 3.15, 2.2, 3.0, 4.6, "\n".join(f"• {x['title'][:48]}" for x in items[:9]), size=10)

    # 13 next steps
    s = d.slide("Recommended next steps", "From dominant finding patterns (recommender stored as data)")
    rows = [["Pattern", "Next package", "Evidence"]] + [[r["pattern"], r["package"], ", ".join(r["evidence"])[:60]]
                                                       for r in result.recommendations]
    d.table(s, 0.5, 1.6, 12.3, rows if len(rows) > 1 else rows + [["—", "", ""]], [4.4, 3.6, 4.3], size=11)

    # 14 method & assumptions
    s = d.slide("Method, evidence and assumptions")
    d.text(s, 0.8, 1.5, 11.8, 5.3,
           f"• Read-only, metadata-only harvest; sealed snapshot {result.snapshot_id} (content hash {ctx.manifest.get('content_hash', '')[:16]}).\n"
           f"• {len(sc['checks'])} checks evaluated from the catalogue ({result.catalogue_hash}); thresholds by rubric reference "
           f"({result.rubric_version}).\n"
           "• Root-cause clustering by query signature, warehouse, dashboard, pipeline and owner.\n"
           "• Value: measured credits × contract price, annualised (p50 and conservative); risk: likelihood × impact.\n"
           "• Confidence derives from evidence completeness; below the review floor a named reviewer must approve.\n"
           "• The same snapshot and rubric version always reproduce these numbers.", size=15)
    d.prs.save(str(path))
    return len(d.prs.slides)
