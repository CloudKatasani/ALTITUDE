"""Board pack (M12): one set of numbers from one sealed snapshot, rendered as PPTX, PDF and an HTML dashboard.

``board_numbers`` is computed once (Package D scorecard and suggestions, Well-Architected rating, AI-readiness)
and every rendering embeds its ``numbers_hash`` — the deck and the dashboard provably read the same snapshot.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from jinja2 import Environment, FileSystemLoader, select_autoescape

from altitude.core.ids import content_hash
from altitude.core.rubric import RubricSet, default_rubrics
from altitude.core.snapshot import Snapshot
from altitude.core.store import json_safe
from altitude.core.theme import render as render_tokens
from altitude.core.theme import token

from . import pdf

_ENV = Environment(loader=FileSystemLoader(str(Path(__file__).parent / "templates")),
                   autoescape=select_autoescape(["html"]))


def board_numbers(snapshot: Snapshot, tier: str = "D2", rubrics: RubricSet | None = None,
                  top_n: int = 5) -> dict[str, Any]:  # threshold-ok: slide row count
    from altitude.airready.scoring import score as ai_score
    from altitude.diagnostics.engine import run as diag_run
    from altitude.foundation.wa import run as wa_run

    rubrics = rubrics or default_rubrics()
    result = diag_run(snapshot, tier=tier, rubrics=rubrics)
    wa = wa_run(snapshot, rubrics=rubrics)
    ai = ai_score(snapshot, rubrics=rubrics)
    sugg = result.suggestions
    value = sum(float((s["value"]["annualised_usd"] or {}).get("p50") or 0) for s in sugg)
    conservative = sum(float((s["value"]["annualised_usd"] or {}).get("conservative") or 0) for s in sugg)
    nums = {
        "snapshot_id": snapshot.snapshot_id, "snapshot_hash": snapshot.content_hash,
        "rubric_version": rubrics.version, "tier": tier,
        "platform_health_index": result.scorecard["platform_health_index"],
        "domains": [{"id": d["id"], "name": d.get("name", d["id"]), "score": d["score"]}
                    for d in result.scorecard["domains"]],
        "pillars": [{"pillar": p["pillar"], "score": p["score"]} for p in result.scorecard["pillars"]],
        "suggestions": len(sugg),
        "quick_wins": sum(1 for s in sugg if s["priority"]["quadrant"] == "quick_win"),
        "identified_value_p50_usd": round(value, 2), "identified_value_conservative_usd": round(conservative, 2),
        "top": [{"id": s["id"], "title": s["title"], "value_p50_usd": (s["value"]["annualised_usd"] or {}).get("p50"),
                 "quadrant": s["priority"]["quadrant"], "evidence_refs": s["evidence_refs"]}
                for s in sorted(sugg, key=lambda s: s["priority"]["rank"])[:top_n]],
        "well_architected": {"overall": wa.overall, "rating": wa.rating, "findings": len(wa.findings),
                             "report_hash": wa.report_hash},
        "ai_readiness": {"overall": ai.overall, "level": ai.maturity_level, "label": ai.maturity_label,
                         "score_hash": ai.score_hash},
        "recommendations": result.recommendations,
    }
    nums = json_safe(nums)
    nums["numbers_hash"] = content_hash(nums, 16)
    return nums


def _usd(v: float | None) -> str:
    return f"${float(v or 0):,.0f}"


def html_dashboard(nums: dict, theme: str | None = None) -> str:
    html = _ENV.get_template("board.html").render(n=nums, usd=_usd, theme=lambda k: token(k, theme))
    return render_tokens(html, theme)


def pdf_pack(nums: dict, theme: str | None = None) -> bytes:
    t = lambda k: token(k, theme)  # noqa: E731
    p1 = [(f"{t('PRODUCT_NAME')} board pack", 22, True), (f"Snapshot {nums['snapshot_id']}", 11, False),
          (f"Numbers hash {nums['numbers_hash']} · rubric {nums['rubric_version']}", 9, False), ("", 8, False),
          (f"{t('HEALTH_INDEX_NAME')}: {nums['platform_health_index']}", 16, True),
          (f"Identified annualised value (p50): {_usd(nums['identified_value_p50_usd'])}  "
           f"(conservative {_usd(nums['identified_value_conservative_usd'])})", 12, False),
          (f"Suggestions: {nums['suggestions']} · quick wins: {nums['quick_wins']}", 12, False),
          (f"Well-Architected: {nums['well_architected']['overall']} ({nums['well_architected']['rating']})", 12, False),
          (f"AI readiness: {nums['ai_readiness']['overall']} (level {nums['ai_readiness']['level']}, "
           f"{nums['ai_readiness']['label']})", 12, False)]
    p2 = [("Top suggestions", 18, True)] + [
        (f"{i}. {s['title']} — {_usd(s['value_p50_usd'])} ({s['quadrant']}) · {s['id']}", 11, False)
        for i, s in enumerate(nums["top"], 1)]
    p2 += [("", 8, False), ("Domain scores (1-5)", 14, True)] + [
        (f"{d['id']} {d['name']}: {d['score']}", 11, False) for d in nums["domains"]]
    return pdf.render([p1, p2], title=f"{t('PRODUCT_NAME')} board pack {nums['snapshot_id']}",
                      creator=t("COMPANY_NAME"))


def pptx_pack(nums: dict, path: Path, theme: str | None = None) -> int:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    t = lambda k: token(k, theme)  # noqa: E731
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # threshold-ok: 16:9 slide size
    prs.core_properties.subject = f"snapshot_id={nums['snapshot_id']}; numbers_hash={nums['numbers_hash']}"
    prs.core_properties.author = t("COMPANY_NAME")

    def slide(title: str, lines: list[str]) -> None:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        box = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1))
        p = box.text_frame.paragraphs[0]
        p.text, p.font.size, p.font.bold = title, Pt(30), True
        p.font.color.rgb = RGBColor.from_string(t("COLOR_PRIMARY").lstrip("#"))
        body = s.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12), Inches(5.5)).text_frame
        body.word_wrap = True
        for i, line in enumerate(lines):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text, para.font.size = line, Pt(18)
        foot = s.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(12), Inches(0.4)).text_frame.paragraphs[0]
        foot.text = f"{nums['snapshot_id']} · numbers {nums['numbers_hash']}"
        foot.font.size = Pt(10)

    slide(f"{t('PRODUCT_NAME')} — board pack", [
        f"{t('HEALTH_INDEX_NAME')}: {nums['platform_health_index']}",
        f"Identified annualised value (p50): {_usd(nums['identified_value_p50_usd'])}",
        f"Suggestions: {nums['suggestions']} (quick wins {nums['quick_wins']})",
        f"Well-Architected: {nums['well_architected']['overall']} ({nums['well_architected']['rating']})",
        f"AI readiness: {nums['ai_readiness']['overall']} (level {nums['ai_readiness']['level']})"])
    slide("Top suggestions", [f"{s['title']} — {_usd(s['value_p50_usd'])} [{s['id']}]" for s in nums["top"]])
    slide("Scorecard by domain", [f"{d['id']} {d['name']}: {d['score']}" for d in nums["domains"]])
    slide("Recommended next steps", [f"{r.get('package', r.get('id', ''))}: {r.get('rationale', r.get('reason', ''))}"
                                     for r in nums["recommendations"]] or ["No follow-on package triggered."])
    prs.save(str(path))
    return len(prs.slides)


def build(snapshot: Snapshot, out_dir: str | Path, tier: str = "D2", theme: str | None = None) -> dict:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    nums = board_numbers(snapshot, tier)
    (out / "board_dashboard.html").write_text(html_dashboard(nums, theme))
    (out / "board_pack.pdf").write_bytes(pdf_pack(nums, theme))
    slides = pptx_pack(nums, out / "board_pack.pptx", theme)
    import json

    (out / "board_numbers.json").write_text(json.dumps(nums, indent=2, sort_keys=True) + "\n")
    return {"numbers_hash": nums["numbers_hash"], "snapshot_id": nums["snapshot_id"], "slides": slides,
            "files": sorted(p.name for p in out.iterdir())}
