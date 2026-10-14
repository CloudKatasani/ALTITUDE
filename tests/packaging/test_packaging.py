import re

import sqlglot
from pptx import Presentation

from altitude.core.generated import is_stamped
from altitude.packaging.board_pack import build
from altitude.packaging.generate import generate
from altitude.packaging.native_app import bundle


def test_native_app_bundle_parses_and_is_least_privilege():
    b = bundle()
    assert set(b) == {"manifest.yml", "setup.sql", "README.md"}
    setup = b["setup.sql"]
    assert not re.search(r"(USE ROLE|GRANT ROLE|TO ROLE)\s+ACCOUNTADMIN", setup.upper())
    assert "EXTERNAL_ACCESS_INTEGRATIONS" not in setup.upper() and "CREATE EXTERNAL ACCESS" not in setup.upper()
    assert "IMPORTED PRIVILEGES ON SNOWFLAKE DB" in b["manifest.yml"]
    views = re.findall(r"CREATE OR REPLACE VIEW DETECTORS\.(DX_[A-Z]+_\d+)", setup)
    assert len(views) >= 33
    # every detector view body is valid Snowflake SQL with no unbound $parameters
    for stmt in re.findall(r"CREATE OR REPLACE VIEW DETECTORS\.\w+ AS (.*?);\n", setup, flags=re.S):
        assert not re.search(r"\$[A-Za-z_]", stmt)  # no unbound $parameters (regex anchors are fine)
        sqlglot.parse_one(stmt, read="snowflake")
    assert all(is_stamped(c) for c in generate().values())


def test_board_pack_deck_pdf_and_dashboard_share_one_snapshot(seeded_snapshot, tmp_path):
    snap, _ = seeded_snapshot
    out = build(snap, tmp_path / "board")
    h = out["numbers_hash"]
    html = (tmp_path / "board" / "board_dashboard.html").read_text()
    pdf = (tmp_path / "board" / "board_pack.pdf").read_bytes()
    deck = Presentation(str(tmp_path / "board" / "board_pack.pptx"))
    assert f'data-numbers-hash="{h}"' in html and f'data-snapshot-id="{snap.snapshot_id}"' in html
    assert pdf.startswith(b"%PDF-1.4") and h.encode() in pdf and pdf.rstrip().endswith(b"%%EOF")
    assert h in deck.core_properties.subject and out["slides"] == 4
    deck_text = " ".join(sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame)
    phi = re.search(r'id="phi">([^<]+)<', html).group(1)
    value = re.search(r'id="value">([^<]+)<', html).group(1)
    assert phi in deck_text and value in deck_text and value.encode() in pdf
    again = build(snap, tmp_path / "board2")
    assert again["numbers_hash"] == h
