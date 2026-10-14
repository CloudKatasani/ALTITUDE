"""Every deliverable from one sealed snapshot; white-label via theme tokens."""

import csv
import io
import json
import re

import jsonschema
import pytest
from openpyxl import load_workbook
from pptx import Presentation

from altitude.core.generated import is_stamped
from altitude.core.paths import repo_root
from altitude.core.theme import tokens
from altitude.diagnostics.exports import ALL_FORMATS, FINDINGS_COLUMNS, export_all, findings_csv, jira_csv
from altitude.diagnostics.report import render_report

SCHEMA = json.loads((repo_root() / "services/altitude/core/schemas/diagnostics_suggestion.json").read_text())


@pytest.fixture(scope="module")
def exported(tmp_path_factory, seeded_result):
    out = tmp_path_factory.mktemp("diag_export")
    return out, export_all(seeded_result, out, formats=ALL_FORMATS)


def test_all_deliverables_are_written(exported):
    out, written = exported
    assert set(written) == set(ALL_FORMATS)
    for fmt, path in written.items():
        assert path.exists(), fmt


def test_suggestions_json_contract_for_observability(exported, seeded_result):
    _, written = exported
    doc = json.loads(written["suggestions_json"].read_text())
    assert doc["snapshot_id"] == seeded_result.snapshot_id and doc["rubric_version"] == seeded_result.rubric_version
    assert doc["kpi_ids"] == sorted({k for s in doc["suggestions"] for k in s["kpis_to_watch"]})
    for s in doc["suggestions"]:
        jsonschema.validate(s, SCHEMA)


def test_findings_register_csv_and_xlsx(exported, seeded_result):
    _, written = exported
    rows = list(csv.DictReader(io.StringIO(written["findings_csv"].read_text())))
    assert list(rows[0]) == FINDINGS_COLUMNS and len(rows) == len(seeded_result.suggestions)
    assert all(r["snapshot_id"] == seeded_result.snapshot_id and r["rubric_version"] == seeded_result.rubric_version
               and r["evidence_refs"] and r["value_method"] and r["confidence"] and r["rollback_ref"] for r in rows)
    wb = load_workbook(written["findings_xlsx"])
    assert wb.sheetnames == ["Findings", "Evidence", "Scorecard", "About"]
    assert wb["Findings"].max_row == len(rows) + 1
    about = {r[0].value: r[1].value for r in wb["About"].iter_rows(min_row=2)}
    assert about["snapshot_id"] == seeded_result.snapshot_id and about["rubric_version"] == seeded_result.rubric_version


def test_backlogs_import_shape(exported, seeded_result):
    _, written = exported
    jira = list(csv.reader(io.StringIO(written["jira_csv"].read_text())))
    assert jira[0][:4] == ["Issue Type", "Summary", "Description", "Priority"]
    assert len(jira) == len(seeded_result.suggestions) + 1
    assert all(seeded_result.snapshot_id in row[2] and seeded_result.rubric_version in row[2] for row in jira[1:])
    ado = list(csv.DictReader(io.StringIO(written["ado_csv"].read_text())))
    assert {r["Work Item Type"] for r in ado} == {"Task"} and all(r["Priority"] in {"1", "2", "3", "4"} for r in ado)
    assert all(f"rubric-{seeded_result.rubric_version}" in r["Tags"] for r in ado)


def test_remediation_pr_pack(exported, seeded_result):
    _, written = exported
    pack = written["pr_pack"]
    manifest = json.loads((pack / "manifest.json").read_text())
    assert manifest["snapshot_id"] == seeded_result.snapshot_id and is_stamped((pack / "manifest.json").read_text())
    assert len(manifest["artifacts"]) == len(seeded_result.suggestions)
    for a in manifest["artifacts"]:
        fwd, rb = pack / a["artifact_ref"], pack / a["rollback_ref"]
        assert is_stamped(fwd.read_text()) and is_stamped(rb.read_text())
        assert "DO NOT EDIT" in fwd.read_text().splitlines()[0]
    readme = (pack / "README.md").read_text()
    assert is_stamped(readme) and tokens()["PRODUCT_NAME"] in readme


def test_executive_readout_deck(exported, seeded_result):
    _, written = exported
    prs = Presentation(str(written["deck_pptx"]))
    assert 12 <= len(prs.slides) <= 15
    text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    t = tokens()
    assert t["PACKAGE_D_NAME"] in text and t["HEALTH_INDEX_NAME"] in text
    assert seeded_result.snapshot_id in text and seeded_result.rubric_version in text
    assert seeded_result.rubric_version in prs.core_properties.subject


def test_report_html_uses_tokens_only_and_white_labels(seeded_result):
    template = (repo_root() / "services/altitude/diagnostics/templates/report.html").read_text()
    assert not re.search(r"#[0-9A-Fa-f]{6}\b", template)
    default = render_report(seeded_result)
    partner = render_report(seeded_result, theme="partner_example")
    assert "{{" not in default and "{{" not in partner
    p = tokens("partner_example")
    assert p["PRODUCT_NAME"] in partner and p["COLOR_PRIMARY"] in partner and p["PACKAGE_D_NAME"] in partner
    assert tokens()["COMPANY_NAME"] not in partner
    for s in seeded_result.suggestions:
        assert s["id"] in default
    for e in seeded_result.evidence:
        assert e["id"] in default


def test_savings_and_risk_model_has_explicit_assumptions(exported, seeded_result):
    _, written = exported
    wb = load_workbook(written["savings_xlsx"])
    assert wb.sheetnames == ["Assumptions", "Savings", "Risk", "Summary", "About"]
    assumptions = {r[0].value: r[1].value for r in wb["Assumptions"].iter_rows(min_row=2)}
    assert assumptions["contract_price_per_credit"] == 3.0
    assert "rubric.diagnostics.checks.cost_01.recoverable_idle_share" in assumptions
    savings = list(wb["Savings"].iter_rows(min_row=2, values_only=True))
    total = sum((s["value"]["annualised_usd"] or {}).get("p50", 0) for s in seeded_result.suggestions)
    assert sum(r[5] for r in savings) == pytest.approx(total, abs=len(savings))
    assert any(str(c.value).startswith("=SUM") for row in wb["Summary"].iter_rows() for c in row)


def test_status_overlay_flows_into_exports(seeded_result):
    recs = [dict(s) for s in seeded_result.suggestions]
    recs[0] = {**recs[0], "status": "accepted", "reviewer": "lead@client"}
    assert "lead@client" in findings_csv(seeded_result, recs)
    assert jira_csv(seeded_result, recs).count("\n") >= len(recs)
